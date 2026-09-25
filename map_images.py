from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


doc = Document("Nội dung bài báo cáo CĐTN.docx")
lines: list[str] = []
for index, paragraph in enumerate(doc.paragraphs):
    embeds = paragraph._p.xpath(".//a:blip/@r:embed")
    if not embeds:
        continue
    targets = [str(doc.part.rels[embed].target_ref) for embed in embeds]
    nearby = []
    for nearby_index in range(max(0, index - 2), min(len(doc.paragraphs), index + 4)):
        text = " ".join(doc.paragraphs[nearby_index].text.split())
        if text:
            nearby.append(f"{nearby_index}:{text}")
    lines.append(f"P{index}: {', '.join(targets)}")
    lines.extend(f"  {item}" for item in nearby)
Path("tmp/source_extracts/image_map.txt").write_text("\n".join(lines), encoding="utf-8")
