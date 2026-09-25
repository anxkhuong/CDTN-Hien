from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from pypdf import PdfReader


def inspect_docx(path: Path) -> dict:
    doc = Document(path)
    paragraphs = []
    for index, paragraph in enumerate(doc.paragraphs):
        text = paragraph.text.strip()
        if text:
            paragraphs.append(
                {
                    "index": index,
                    "style": paragraph.style.name if paragraph.style else "",
                    "text": text,
                }
            )
    tables = []
    for table_index, table in enumerate(doc.tables):
        rows = []
        for row in table.rows:
            rows.append([cell.text.strip() for cell in row.cells])
        tables.append({"index": table_index, "rows": rows})
    sections = []
    for section in doc.sections:
        sections.append(
            {
                "page_width_cm": round(section.page_width.cm, 2),
                "page_height_cm": round(section.page_height.cm, 2),
                "top_margin_cm": round(section.top_margin.cm, 2),
                "bottom_margin_cm": round(section.bottom_margin.cm, 2),
                "left_margin_cm": round(section.left_margin.cm, 2),
                "right_margin_cm": round(section.right_margin.cm, 2),
                "header_distance_cm": round(section.header_distance.cm, 2),
                "footer_distance_cm": round(section.footer_distance.cm, 2),
            }
        )
    drawings = len(doc._element.xpath(".//w:drawing"))
    pictures = len(doc._element.xpath(".//pic:pic"))
    explicit_page_breaks = len(doc._element.xpath(".//w:br[@w:type='page']"))
    relationships = []
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            relationships.append(str(rel.target_ref))
    return {
        "type": "docx",
        "paragraph_count": len(doc.paragraphs),
        "nonempty_paragraph_count": len(paragraphs),
        "table_count": len(doc.tables),
        "sections": sections,
        "drawing_count": drawings,
        "picture_count": pictures,
        "explicit_page_break_count": explicit_page_breaks,
        "image_relationships": relationships,
        "paragraphs": paragraphs,
        "tables": tables,
    }


def inspect_pdf(path: Path) -> dict:
    reader = PdfReader(path)
    pages = []
    for page_number, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        pages.append({"page": page_number, "text": text.strip()})
    return {"type": "pdf", "page_count": len(reader.pages), "pages": pages}


def main() -> None:
    output_dir = Path("tmp/source_extracts")
    output_dir.mkdir(parents=True, exist_ok=True)
    for raw_path in sys.argv[1:]:
        path = Path(raw_path)
        data = inspect_docx(path) if path.suffix.lower() == ".docx" else inspect_pdf(path)
        output_path = output_dir / f"{path.stem}.json"
        output_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
        if data["type"] == "docx":
            outline_path = output_dir / f"{path.stem}.outline.txt"
            outline_lines = []
            for item in data["paragraphs"]:
                text_value = item["text"].replace("\n", " / ")
                if len(text_value) <= 150:
                    outline_lines.append(
                        f"{item['index']:04d}\t{item['style']}\t{text_value}"
                    )
            outline_path.write_text("\n".join(outline_lines), encoding="utf-8")
        print(f"processed {data['type']}: {output_path.as_posix().encode('ascii', 'backslashreplace').decode('ascii')}")


if __name__ == "__main__":
    main()
