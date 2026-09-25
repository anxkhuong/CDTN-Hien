from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parent
FONT = "Times New Roman"


PAGES = [
    {
        "title": "2.5. Cơ sở lý thuyết của phát hiện bất thường ngoại quan",
        "paragraphs": [
            "Phát hiện bất thường ngoại quan là bài toán xác định những sai lệch có thể quan sát được trên ảnh sản phẩm so với trạng thái được xem là bình thường. Trong kiểm tra công nghiệp, sai lệch có thể chỉ chiếm một vùng rất nhỏ, không xuất hiện lặp lại và thay đổi đáng kể về hình dạng, màu sắc hoặc mức độ tương phản. Vì vậy, bài toán không chỉ là nhận dạng một lớp đối tượng đã biết, mà còn phải mô hình hóa miền bình thường đủ chặt để phát hiện những biểu hiện chưa từng xuất hiện trong dữ liệu huấn luyện. Cách tiếp cận này đặc biệt phù hợp khi doanh nghiệp có nhiều ảnh đạt chuẩn nhưng rất ít ảnh lỗi được gán nhãn đầy đủ.",
            "Một hệ thống kiểm tra ảnh thường thực hiện hai nhiệm vụ liên quan nhưng không đồng nhất. Nhiệm vụ cấp ảnh trả lời ảnh đầu vào có bất thường hay không; nhiệm vụ định vị tạo bản đồ điểm hoặc mặt nạ để chỉ ra vùng nghi ngờ. Điểm cấp ảnh có thể được tổng hợp từ điểm lớn nhất hoặc một nhóm điểm cao trên bản đồ bất thường, trong khi mặt nạ cuối cùng còn phụ thuộc vào ngưỡng, giới hạn vùng đối tượng và hậu xử lý. Vì thế, một mô hình phân loại tốt chưa chắc tạo được vùng lỗi chính xác, và một bản đồ điểm giàu thông tin vẫn có thể cho kết quả nhị phân kém nếu lựa chọn ngưỡng không phù hợp.",
            "Cơ sở lý thuyết trong mục này làm rõ chuỗi lập luận từ biểu diễn ảnh, trích xuất đặc trưng, mô hình hóa phân bố bình thường, đo khoảng cách trong không gian nhúng đến đánh giá và hiệu chỉnh đầu ra. Trọng tâm được đặt vào các phương pháp học một lớp và không giám sát vì chúng phù hợp với điều kiện dữ liệu của chuyên đề. Các khái niệm được liên hệ trực tiếp với PatchCore [4], DRAEM [5], Isolation Forest [2] và các lựa chọn tiền xử lý, hậu xử lý được sử dụng trong quy trình nghiên cứu."
        ],
    },
    {
        "title": "2.5.1. Kiểm tra chất lượng bằng thị giác máy tính",
        "paragraphs": [
            "Kiểm tra ngoại quan truyền thống dựa nhiều vào kinh nghiệm của nhân viên và thường chịu ảnh hưởng bởi mệt mỏi, tốc độ dây chuyền, điều kiện chiếu sáng và tiêu chí đánh giá giữa các ca làm việc. Thị giác máy tính không loại bỏ vai trò của con người mà chuyển phần quan sát lặp lại thành một quy trình có thể đo lường. Camera, môi trường chụp và thuật toán tạo thành một hệ đo; nếu một thành phần thay đổi, phân bố ảnh cũng thay đổi. Do đó, chất lượng của hệ thống phụ thuộc đồng thời vào dữ liệu thu nhận và mô hình, không thể chỉ xem độ chính xác của thuật toán là tiêu chí duy nhất.",
            "Trong thực tế, khái niệm “lỗi” được xác định bởi yêu cầu chất lượng của từng mã sản phẩm. Cùng một vết xước có thể được chấp nhận ở vùng khuất nhưng bị loại ở mặt hiển thị; một sai lệch màu nhỏ có thể quan trọng với bao bì nhưng ít ý nghĩa với chi tiết cơ khí. Vì vậy, trước khi xây dựng mô hình cần xác định đơn vị kiểm tra, vùng quan tâm, mức độ lỗi và hành động sau cảnh báo. Cấu trúc này giúp phân biệt bất thường thị giác với lỗi chất lượng thực sự, đồng thời tạo cơ sở cho việc đặt ngưỡng theo chi phí sai sót.",
            "Đối với chuyên đề, ảnh được xử lý theo từng mẫu sản phẩm để tránh trộn các phân bố bình thường khác nhau. Mô hình tạo điểm bất thường và vùng nghi ngờ, còn quyết định cuối cùng được kiểm soát bằng ngưỡng và giới hạn tiền cảnh. Thiết kế này phù hợp với nguyên tắc truy vết: mỗi cảnh báo phải gắn với ảnh gốc, mã sản phẩm, điểm số và vùng được đánh dấu. Khi triển khai, dữ liệu xác nhận của nhân viên kiểm tra có thể được lưu lại để hiệu chỉnh ngưỡng và theo dõi sự thay đổi của điều kiện sản xuất theo thời gian."
        ],
        "table": [
            ["Thành phần", "Vai trò", "Rủi ro cần kiểm soát"],
            ["Thu nhận ảnh", "Tạo đầu vào ổn định", "Ánh sáng, góc chụp, độ rung"],
            ["Mô hình", "Tạo điểm và bản đồ bất thường", "Lệch miền, đặc trưng không phù hợp"],
            ["Quy tắc quyết định", "Chuyển điểm thành cảnh báo", "Ngưỡng và chi phí sai sót"],
        ],
        "caption": "Bảng 2.2. Các thành phần của hệ thống kiểm tra ngoại quan.",
    },
    {
        "title": "2.5.2. Phân loại khuyết tật và mức độ biểu hiện",
        "paragraphs": [
            "Khuyết tật ngoại quan có thể được mô tả theo nguồn gốc hoặc theo biểu hiện trên ảnh. Theo biểu hiện, nhóm lỗi hình học gồm thiếu chi tiết, thừa chi tiết, lệch vị trí, biến dạng và sai kích thước; nhóm lỗi bề mặt gồm xước, lõm, nứt, bẩn, bong tróc hoặc đốm màu; nhóm lỗi cấu trúc gồm lắp sai quan hệ giữa các bộ phận. Cách phân loại này có ý nghĩa đối với lựa chọn đặc trưng. Lỗi bề mặt thường cần đặc trưng cục bộ ở độ phân giải cao, trong khi lỗi cấu trúc đòi hỏi trường tiếp nhận lớn hơn và khả năng biểu diễn quan hệ không gian.",
            "Một đặc điểm quan trọng là độ khó không tỷ lệ tuyến tính với kích thước lỗi. Vùng lỗi lớn nhưng tương phản rõ có thể dễ phát hiện, trong khi một vết xước mảnh nằm trên họa tiết phức tạp dễ bị hòa vào nền đặc trưng. Độ khó còn phụ thuộc tính hiếm, độ biến thiên của mẫu bình thường và mức tương đồng giữa lỗi với các biến đổi hợp lệ. Các phản xạ sáng, nếp gấp cho phép hoặc sai khác vị trí nhỏ có thể tạo điểm bất thường cao dù không phải lỗi, dẫn đến dương tính giả.",
            "Do dữ liệu lỗi thường không bao phủ đầy đủ mọi tình huống, mô hình học một lớp tập trung mô tả phạm vi biến thiên bình thường. Khi một đặc trưng kiểm tra nằm xa vùng đã quan sát, hệ thống tăng điểm bất thường. Cách diễn giải này không khẳng định nguyên nhân lỗi mà chỉ cho biết mức không phù hợp với dữ liệu tham chiếu. Vì vậy, kết quả nên được trình bày dưới dạng hỗ trợ quyết định, kèm vùng nghi ngờ và mức tin cậy, thay vì coi mọi điểm vượt ngưỡng đều là một loại lỗi đã được xác định."
        ],
        "table": [
            ["Nhóm biểu hiện", "Ví dụ", "Yêu cầu biểu diễn"],
            ["Bề mặt", "Xước, bẩn, bong tróc", "Chi tiết cục bộ, độ phân giải cao"],
            ["Hình học", "Thiếu, thừa, lệch vị trí", "Hình dạng và bố cục"],
            ["Cấu trúc", "Sai quan hệ giữa bộ phận", "Ngữ cảnh và trường tiếp nhận lớn"],
        ],
        "caption": "Bảng 2.3. Phân nhóm khuyết tật theo biểu hiện trên ảnh.",
    },
    {
        "title": "2.5.3. Biểu diễn ảnh số và không gian màu",
        "paragraphs": [
            "Ảnh số màu thường được biểu diễn bởi tensor X ∈ R^(H×W×C), trong đó H và W là chiều cao, chiều rộng, còn C là số kênh. Với ảnh RGB, mỗi pixel chứa ba thành phần đỏ, lục và lam. Giá trị pixel chỉ là phép đo cường độ tại vị trí rời rạc, chưa mang sẵn ý nghĩa về vật thể hay khuyết tật. Ý nghĩa được hình thành qua quan hệ giữa các pixel lân cận và qua đặc trưng học được ở nhiều mức. Khi ảnh có kênh alpha, alpha mô tả độ trong suốt hoặc vùng tiền cảnh; sử dụng đúng kênh này có thể tách đối tượng chính xác hơn phép ngưỡng màu đơn giản.",
            "Không gian RGB thuận tiện cho mạng đã được tiền huấn luyện trên ảnh tự nhiên, nhưng độ sáng và màu sắc bị trộn trong ba kênh. Các không gian HSV hoặc Lab tách tương đối thành phần sáng khỏi sắc độ, hữu ích khi cần phân tích sai lệch màu hoặc xây dựng quy tắc tiền xử lý. Tuy nhiên, đổi không gian màu không tự động cải thiện phát hiện bất thường; phép biến đổi có thể khuếch đại nhiễu ở vùng tối hoặc tạo gián đoạn quanh biên góc màu. Lựa chọn cần dựa trên loại lỗi và đặc tính thu nhận ảnh.",
            "Mạng tích chập thường nhận tensor dấu phẩy động đã được đưa về một miền giá trị thống nhất. Với bộ trích xuất tiền huấn luyện trên ImageNet [17], ảnh thường được chuẩn hóa theo trung bình và độ lệch chuẩn của dữ liệu tiền huấn luyện. Việc giữ đúng thứ tự kênh, phạm vi giá trị và tham số chuẩn hóa là điều kiện quan trọng; sai khác nhỏ ở bước này có thể làm toàn bộ không gian nhúng thay đổi. Trong chuyên đề, tính nhất quán của chuỗi đọc ảnh, xử lý alpha, đổi kích thước và chuẩn hóa được xem như một phần của mô hình."
        ],
    },
    {
        "title": "2.5.4. Độ sáng, tương phản và biến thiên điều kiện chụp",
        "paragraphs": [
            "Độ sáng biểu thị mức cường độ tổng thể, còn tương phản mô tả mức chênh lệch giữa vùng sáng và tối. Trong ảnh công nghiệp, cùng một bề mặt có thể tạo giá trị pixel khác nhau do công suất đèn, vị trí nguồn sáng, độ bóng vật liệu và góc quan sát. Những biến đổi này không nhất thiết là lỗi nhưng có thể làm đặc trưng rời xa ngân hàng tham chiếu. Vì vậy, dữ liệu huấn luyện bình thường phải bao phủ được phạm vi vận hành hợp lệ thay vì chỉ gồm các ảnh gần như giống hệt nhau.",
            "Vật thể tối đặt trên nền tối là trường hợp đặc biệt khó đối với phép tạo mặt nạ dựa trên ngưỡng cường độ. Nếu tiêu chí tiền cảnh là pixel đủ sáng, một phần hợp lệ của đối tượng có thể bị loại bỏ. Sai lệch mặt nạ sau đó lan truyền đến tất cả bước sau: đặc trưng không được tính trên vùng đã mất, bản đồ bất thường bị cắt và điểm cấp ảnh có thể giảm giả tạo. Đây là ví dụ cho thấy tiền xử lý không trung lập; nó chứa giả định về dữ liệu và cần được đánh giá định lượng trên toàn bộ các mã sản phẩm.",
            "Các biện pháp ổn định gồm cố định camera và nguồn sáng, sử dụng hộp chụp, khóa phơi sáng và cân bằng trắng, đồng thời lưu thông số thiết bị cùng ảnh. Chuẩn hóa histogram hoặc tăng cường độ sáng có thể giảm một phần biến thiên nhưng cũng có nguy cơ làm thay đổi tín hiệu lỗi. Một nguyên tắc an toàn là chỉ áp dụng phép biến đổi mà ảnh hưởng của nó có thể kiểm chứng trên cả ảnh bình thường và ảnh lỗi. Khi điều kiện chụp thay đổi ngoài phạm vi đã biết, hệ thống nên phát hiện lệch miền và yêu cầu hiệu chuẩn lại thay vì tiếp tục dùng ngưỡng cũ."
        ],
    },
    {
        "title": "2.5.5. Tiền xử lý, đổi kích thước và chuẩn hóa",
        "paragraphs": [
            "Tiền xử lý biến ảnh thu nhận thành đầu vào có kích thước và phân bố phù hợp với bộ trích xuất đặc trưng. Đổi kích thước là thao tác cần thiết nhưng tạo ra đánh đổi: độ phân giải thấp giảm chi phí tính toán và dung lượng ngân hàng bộ nhớ, trong khi độ phân giải cao giữ được vết lỗi nhỏ. Nếu tỷ lệ khung hình bị thay đổi, hình dạng đối tượng và cấu trúc bề mặt cũng bị biến dạng. Có thể giữ tỷ lệ rồi đệm biên, hoặc đổi trực tiếp khi dữ liệu đã có bố cục ổn định; lựa chọn phải giống nhau ở giai đoạn xây dựng ngân hàng và suy luận.",
            "Nội suy gần nhất phù hợp với mặt nạ nhị phân vì không tạo giá trị lớp trung gian. Đối với ảnh màu, nội suy song tuyến tính hoặc hai khối thường cho biên mượt hơn nhưng có thể làm mờ khuyết tật rất mảnh. Chuẩn hóa giá trị pixel giúp các kênh có thang đo phù hợp với mô hình tiền huấn luyện. Ngoài ra, tách nền có thể giảm nhiễu do phông chụp, song chỉ có lợi khi mặt nạ đối tượng đáng tin cậy. Nếu nền đã ổn định, việc xóa nền đôi khi loại bỏ cả thông tin biên hữu ích.",
            "Tăng cường dữ liệu như lật, xoay, thay đổi sáng hoặc thêm nhiễu phải phản ánh biến thiên hợp lệ của quá trình chụp. Một phép xoay không thực tế có thể khiến mô hình coi bố cục sai là bình thường. Với phương pháp không huấn luyện bộ trích xuất như PatchCore, tăng cường còn làm ngân hàng đặc trưng lớn hơn và thay đổi mật độ mẫu. Do đó, tiền xử lý nên được coi là một tập giả thuyết có thể kiểm thử bằng thí nghiệm loại bỏ từng thành phần, thay vì một chuỗi thao tác mặc định."
        ],
        "table": [
            ["Thao tác", "Lợi ích dự kiến", "Rủi ro"],
            ["Đổi kích thước", "Đồng nhất đầu vào", "Mất lỗi nhỏ hoặc méo hình"],
            ["Chuẩn hóa", "Phù hợp mô hình tiền huấn luyện", "Sai thứ tự kênh/tham số"],
            ["Tách nền", "Giảm nhiễu ngoài vật thể", "Cắt mất vùng vật thể tối"],
            ["Tăng cường", "Bao phủ biến thiên hợp lệ", "Đưa giả định không thực tế"],
        ],
        "caption": "Bảng 2.4. Lợi ích và rủi ro của các bước tiền xử lý.",
    },
    {
        "title": "2.5.6. Mạng nơ-ron tích chập và phép tích chập",
        "paragraphs": [
            "Mạng nơ-ron tích chập (CNN) khai thác tính cục bộ và chia sẻ trọng số để học mẫu hình trên ảnh. Một bộ lọc nhỏ trượt trên không gian ảnh, tính tổ hợp tuyến tính của vùng lân cận rồi tạo bản đồ đặc trưng. Với đầu vào x, trọng số w và độ lệch b, giá trị tại vị trí (i,j) có thể viết y(i,j)=b+Σ_mΣ_uΣ_v w_m(u,v)x_m(i+u,j+v). Cùng một bộ lọc được sử dụng tại nhiều vị trí, nhờ đó số tham số thấp hơn lớp kết nối đầy đủ và đặc trưng có tính tương đương theo dịch chuyển.",
            "Các lớp đầu thường phản ứng với cạnh, góc và kết cấu đơn giản; lớp sâu hơn kết hợp chúng thành họa tiết, bộ phận và cấu trúc ngữ nghĩa. Đối với phát hiện bất thường công nghiệp, đặc trưng quá nông nhạy với nhiễu pixel, trong khi đặc trưng quá sâu có thể bỏ qua vết lỗi nhỏ vì ưu tiên ý nghĩa lớp đối tượng. Vì vậy, nhiều phương pháp ghép đặc trưng từ các tầng trung gian để cân bằng chi tiết không gian và khả năng phân biệt.",
            "Bước nhảy (stride), phần đệm (padding) và kích thước kernel quyết định kích thước bản đồ đầu ra. Stride lớn giảm chi phí nhưng làm lưới đặc trưng thưa hơn, ảnh hưởng trực tiếp đến độ chính xác định vị. Padding giữ kích thước không gian nhưng tạo vùng biên có thông tin nhân tạo. Trong PatchCore, mỗi vị trí trên bản đồ đặc trưng được xem như một patch nhúng; bởi vậy, mọi quyết định làm thay đổi mật độ bản đồ cũng làm thay đổi số phần tử ngân hàng bộ nhớ và độ mịn của bản đồ bất thường."
        ],
    },
    {
        "title": "2.5.7. Hàm kích hoạt, gộp mẫu và trường tiếp nhận",
        "paragraphs": [
            "Sau phép tích chập, hàm kích hoạt đưa tính phi tuyến vào mạng. ReLU được định nghĩa f(z)=max(0,z), đơn giản và giúp giảm hiện tượng gradient suy giảm trong nhiều kiến trúc. Các biến thể như Leaky ReLU giữ một độ dốc nhỏ ở miền âm. Trong mô hình dùng làm bộ trích xuất cố định, lựa chọn hàm kích hoạt đã được xác lập từ quá trình tiền huấn luyện; người sử dụng chủ yếu quyết định lấy đầu ra ở tầng nào và chuẩn hóa đặc trưng ra sao.",
            "Gộp cực đại hoặc gộp trung bình giảm kích thước không gian và tăng tính bền với dịch chuyển nhỏ. Tuy nhiên, gộp mẫu cũng làm mất vị trí chính xác của tín hiệu. Với bài toán phân loại, mất mát này có thể chấp nhận được; với định vị vết xước vài pixel, nó tạo giới hạn độ phân giải. Bản đồ điểm thô thường phải được nội suy trở lại kích thước ảnh, nên đường biên vùng lỗi không thể chi tiết hơn thông tin có trong lưới đặc trưng ban đầu.",
            "Trường tiếp nhận của một phần tử đặc trưng là vùng ảnh đầu vào có thể ảnh hưởng đến phần tử đó. Khi mạng sâu hơn, trường tiếp nhận lý thuyết tăng lên, cho phép mô tả ngữ cảnh rộng. Patch quá nhỏ dễ nhầm một hoa văn hợp lệ với lỗi; patch quá lớn có thể pha loãng tín hiệu cục bộ. Ghép nhiều tầng là cách dùng đồng thời các trường tiếp nhận khác nhau. Trong hệ thống thực tế, lựa chọn tầng và độ phân giải đầu vào cần được đánh giá cùng nhau vì tăng kích thước ảnh không nhất thiết tăng tương ứng độ chi tiết nếu tầng đặc trưng vẫn có stride lớn."
        ],
    },
    {
        "title": "2.5.8. Biểu diễn phân cấp và đặc trưng theo patch",
        "paragraphs": [
            "Biểu diễn phân cấp là cơ sở để biến ảnh pixel thành các vector đặc trưng có ý nghĩa. Ở mỗi tầng, một vị trí không còn mô tả một pixel đơn lẻ mà đại diện cho một vùng ảnh. Có thể xem vector này là đặc trưng patch. Khi các vector từ ảnh bình thường được lưu lại, chúng tạo thành tập tham chiếu của những cấu trúc cục bộ được chấp nhận. Patch ở ảnh kiểm tra được so sánh với tập tham chiếu để xác định mức độ mới lạ.",
            "Đặc trưng từ nhiều tầng thường có số kênh và kích thước không gian khác nhau. Trước khi ghép, bản đồ phải được căn chỉnh bằng nội suy hoặc phép gộp cục bộ. Nếu chỉ nội suy đơn giản, một vector có thể đại diện không nhất quán cho các vùng có trường tiếp nhận khác nhau; vì vậy PatchCore sử dụng cơ chế tổng hợp lân cận để giữ thông tin patch [4]. Sau khi ghép, chuẩn hóa vector giúp khoảng cách ít bị chi phối bởi tầng có biên độ lớn.",
            "Lợi thế của đặc trưng patch là khả năng định vị: điểm bất thường được gắn với vị trí không gian trước khi tổng hợp thành điểm ảnh. Hạn chế là số lượng vector tăng theo số ảnh, số tầng và độ phân giải. Với N ảnh và lưới h×w, ngân hàng có thể chứa xấp xỉ N·h·w vector trước khi rút gọn. Đây là lý do các kỹ thuật coreset có vai trò quan trọng. Ngoài dung lượng, tập patch bình thường cũng cần đủ đa dạng; một vùng hợp lệ nhưng hiếm có thể thiếu láng giềng gần và bị báo lỗi."
        ],
    },
    {
        "title": "2.5.9. Mạng dư ResNet",
        "paragraphs": [
            "ResNet được đề xuất để huấn luyện mạng rất sâu bằng các kết nối tắt [8]. Thay vì buộc một khối học trực tiếp ánh xạ H(x), khối học phần dư F(x)=H(x)−x và tạo đầu ra y=F(x,{W_i})+x. Đường đồng nhất giúp gradient truyền qua nhiều lớp thuận lợi hơn, giảm khó khăn tối ưu khi độ sâu tăng. ResNet vì thế trở thành một họ kiến trúc phổ biến và thường được dùng làm bộ trích xuất đặc trưng trong phát hiện bất thường.",
            "Trong mô hình tiền huấn luyện, các tầng đầu của ResNet giữ cấu trúc biên và kết cấu, các tầng giữa mô tả bộ phận, còn tầng cuối hướng nhiều hơn tới phân loại ImageNet. PatchCore thường lấy đặc trưng trung gian từ backbone Wide ResNet để duy trì độ phân giải không gian và tăng năng lực biểu diễn [4]. Các trọng số có thể được đóng băng, nghĩa là chuyên đề không tối ưu lại hàng triệu tham số từ tập dữ liệu nhỏ. Điều này làm giảm nguy cơ quá khớp và chi phí huấn luyện.",
            "Tuy nhiên, backbone mạnh không bảo đảm phù hợp tuyệt đối với ảnh công nghiệp. ImageNet chứa ảnh tự nhiên với đối tượng và bối cảnh đa dạng, trong khi khuyết tật công nghiệp có thể là thay đổi vật liệu rất nhỏ. Đặc trưng tiền huấn luyện vẫn hữu ích nhờ tính khái quát, nhưng cần lựa chọn tầng, độ phân giải và phương pháp đo khoảng cách thích hợp. Khi miền dữ liệu khác biệt mạnh, có thể cân nhắc tự giám sát hoặc tinh chỉnh có kiểm soát, song phải đánh đổi với nhu cầu dữ liệu và rủi ro làm mất tính ổn định."
        ],
    },
    {
        "title": "2.5.10. Học chuyển giao và bộ dữ liệu tiền huấn luyện",
        "paragraphs": [
            "Học chuyển giao tái sử dụng tri thức từ một nhiệm vụ nguồn cho nhiệm vụ đích. AlexNet [7] và các kiến trúc CNN sau đó chứng minh khả năng học đặc trưng quy mô lớn trên ImageNet [17]. Khi dùng backbone tiền huấn luyện, mô hình phát hiện bất thường thừa hưởng các bộ lọc cạnh, kết cấu và hình dạng thay vì học từ đầu bằng vài chục ảnh bình thường. Đây là một trong những lý do phương pháp dựa trên đặc trưng có thể hoạt động trong chế độ dữ liệu khan hiếm.",
            "Có ba mức sử dụng phổ biến: đóng băng hoàn toàn backbone; chỉ tinh chỉnh các tầng cuối; hoặc tinh chỉnh toàn bộ mạng. Đóng băng cho kết quả ổn định và ít tham số nhưng có thể giữ đặc trưng chưa phù hợp miền. Tinh chỉnh tăng khả năng thích nghi nhưng cần tiêu chí học và dữ liệu đủ lớn. Trong phát hiện bất thường chỉ có ảnh bình thường, nếu tối ưu không cẩn thận, mô hình có thể học ánh xạ suy biến hoặc quá khớp vào bố cục ảnh huấn luyện.",
            "Sự khác biệt giữa miền nguồn và miền đích được gọi là độ lệch miền. Nó xuất hiện do loại vật liệu, thiết bị chụp, thang màu hoặc cấu trúc ảnh. Độ lệch có thể được đánh giá thông qua phân bố điểm của ảnh bình thường mới, khoảng cách đặc trưng và tỷ lệ cảnh báo theo thời gian. Với chuyên đề, sử dụng backbone cố định giúp tách ảnh hưởng của đặc trưng tiền huấn luyện khỏi các thành phần tiền xử lý và hậu xử lý. Khi mở rộng triển khai, việc cập nhật ngân hàng đặc trưng nên được quản lý phiên bản để tránh trộn dữ liệu từ các điều kiện chụp không tương thích."
        ],
    },
    {
        "title": "2.5.11. Khái niệm bất thường và học một lớp",
        "paragraphs": [
            "Bất thường là quan sát có xác suất thấp hoặc không phù hợp với quy luật của dữ liệu tham chiếu. Định nghĩa này phụ thuộc ngữ cảnh: một màu sắc hiếm có thể là lỗi đối với một mã sản phẩm nhưng hoàn toàn bình thường với mã khác. Trong học một lớp, thuật toán chủ yếu quan sát lớp bình thường và cố gắng xác định miền hỗ trợ của lớp đó. Deep SVDD là ví dụ học biểu diễn sao cho ảnh bình thường tập trung quanh một tâm trong không gian đặc trưng [9]. Quan sát xa tâm hoặc xa miền hỗ trợ nhận điểm bất thường cao.",
            "Có thể phân biệt bất thường điểm, bất thường cục bộ và bất thường logic. Bất thường điểm xuất hiện khi toàn bộ mẫu khác biệt; bất thường cục bộ chỉ nằm ở một vùng; bất thường logic xảy ra khi các thành phần riêng lẻ hợp lệ nhưng quan hệ giữa chúng sai. MVTec AD tập trung mạnh vào lỗi bề mặt và cấu trúc cục bộ [1], còn MVTec LOCO AD nhấn mạnh ràng buộc logic [14]. Sự phân biệt này cho thấy mô hình patch cục bộ có thể bỏ sót sai quan hệ ở phạm vi toàn ảnh.",
            "Điểm bất thường không phải xác suất đã được hiệu chuẩn. Nó thường là khoảng cách, năng lượng hoặc sai số tái cấu trúc và chỉ có ý nghĩa tương đối trong cùng cấu hình. Nếu đổi backbone, độ phân giải hay tỷ lệ coreset, thang điểm có thể thay đổi. Vì vậy, ngưỡng phải được xác định lại trên tập hiệu chỉnh tách biệt. Khi không có ảnh lỗi đại diện, có thể chọn ngưỡng theo phân vị của điểm bình thường, nhưng cần ghi rõ đây là kiểm soát tỷ lệ cảnh báo trên dữ liệu tham chiếu chứ không bảo đảm độ nhạy với mọi loại lỗi."
        ],
    },
    {
        "title": "2.5.12. Học có giám sát, không giám sát và bán giám sát",
        "paragraphs": [
            "Học có giám sát sử dụng nhãn lớp hoặc mặt nạ lỗi để tối ưu trực tiếp mục tiêu phân loại, phát hiện hay phân đoạn. Khi bộ nhãn lớn và đại diện, mô hình có thể học ranh giới phù hợp với loại lỗi cụ thể. Hạn chế là lỗi mới không có trong dữ liệu dễ bị bỏ sót, còn chi phí chú thích pixel rất cao. Trong môi trường sản xuất, phân bố lỗi thay đổi và nhiều loại lỗi chỉ xuất hiện sau thời gian dài, khiến việc xây dựng tập huấn luyện cân bằng trở nên khó khăn.",
            "Trong tài liệu phát hiện bất thường công nghiệp, thuật ngữ “không giám sát” thường chỉ tình huống huấn luyện bằng ảnh bình thường mà không dùng nhãn lỗi, dù bản thân việc biết tập huấn luyện là bình thường đã là một dạng thông tin. “Bán giám sát” có thể chỉ sử dụng chủ yếu ảnh bình thường cùng một lượng nhỏ ảnh lỗi, hoặc sử dụng nhãn ở một phần dữ liệu. Vì cách gọi khác nhau giữa các công trình, báo cáo cần mô tả rõ dữ liệu nào được dùng và ở giai đoạn nào thay vì chỉ dựa vào nhãn phương pháp.",
            "Chuyên đề lựa chọn chế độ ảnh bình thường cho huấn luyện, ảnh bình thường và ảnh lỗi cho đánh giá. Cấu trúc này phản ánh bài toán mở: mô hình không học danh sách lỗi đóng mà so sánh mẫu kiểm tra với kiến thức bình thường. Ảnh lỗi trong tập kiểm tra không được đưa vào ngân hàng đặc trưng. Nếu dùng ảnh lỗi để chọn ngưỡng hoặc cấu hình, phần dữ liệu đó trở thành tập xác thực và phải tách khỏi tập báo cáo cuối cùng để tránh đánh giá lạc quan."
        ],
    },
    {
        "title": "2.5.13. Không gian nhúng, khoảng cách và láng giềng gần",
        "paragraphs": [
            "Không gian nhúng biến một patch ảnh thành vector f(p)∈R^d. Hai patch có cấu trúc tương tự được kỳ vọng nằm gần nhau, còn patch khác biệt nằm xa. Khoảng cách Euclid d(x,y)=||x−y||₂ là lựa chọn phổ biến. Nếu vector được chuẩn hóa L2, khoảng cách Euclid có quan hệ đơn điệu với độ tương đồng cosine. Chuẩn hóa giúp giảm ảnh hưởng của độ lớn vector, nhưng có thể loại bỏ thông tin biên độ hữu ích; quyết định này phụ thuộc đặc trưng và phương pháp tham chiếu.",
            "Với ngân hàng M của patch bình thường, điểm cơ bản là s(p)=min_(m∈M)||f(p)−m||₂. Đây là quy tắc một láng giềng gần nhất: chỉ cần một mẫu bình thường tương tự để giảm điểm. Dùng trung bình k láng giềng làm kết quả bền hơn với nhiễu nhưng có thể tăng điểm ở cấu trúc bình thường hiếm. Khi d lớn, khoảng cách có xu hướng tập trung và tìm kiếm chính xác tốn chi phí; rút gọn chiều, coreset và thư viện tìm kiếm gần đúng được dùng để cân bằng tốc độ với độ chính xác.",
            "Khoảng cách chỉ đáng tin khi các chiều đặc trưng có thang đo hợp lý. PaDiM dùng khoảng cách Mahalanobis d_M(x)=√((x−μ)^TΣ⁻¹(x−μ)), xét cả phương sai và tương quan tại từng vị trí [10]. Nếu một chiều biến thiên nhiều trong ảnh bình thường, sai lệch trên chiều đó bị phạt ít hơn. Ngược lại, PatchCore giữ các mẫu cụ thể trong ngân hàng và dùng khoảng cách gần nhất, phù hợp khi phân bố bình thường đa đỉnh và khó mô tả bằng một Gaussian duy nhất."
        ],
    },
    {
        "title": "2.5.14. Isolation Forest và đường cơ sở thống kê",
        "paragraphs": [
            "Isolation Forest phát hiện bất thường bằng cách cô lập quan sát thay vì ước lượng trực tiếp mật độ [2]. Mỗi cây chọn ngẫu nhiên một thuộc tính và một điểm cắt giữa giá trị nhỏ nhất với lớn nhất. Quan sát khác biệt thường bị tách khỏi phần còn lại sau ít bước hơn, nên có độ dài đường đi trung bình ngắn. Điểm bất thường được suy ra từ độ dài này qua chuẩn hóa theo kích thước mẫu. Phương pháp có chi phí hợp lý và không đòi hỏi giả định phân bố Gaussian.",
            "Isolation Forest hoạt động trên vector đặc trưng cố định. Với ảnh, không nên đưa toàn bộ pixel thô có chiều rất cao vào mô hình mà thường cần trích xuất đặc trưng toàn cục hoặc các thống kê như màu, độ sáng, kết cấu. Khi dùng đặc trưng toàn cục, phương pháp phù hợp hơn với phát hiện ảnh khác biệt tổng thể và khó cung cấp mặt nạ lỗi chính xác. Có thể áp dụng trên patch, nhưng số mô hình hoặc số mẫu tăng mạnh và quan hệ không gian không được mô hình hóa trực tiếp.",
            "Trong nghiên cứu, đường cơ sở đơn giản có giá trị vì cho biết cải thiện có đến từ mô hình phức tạp hay chỉ từ dữ liệu dễ phân tách. Một so sánh công bằng cần dùng cùng cách chia dữ liệu, không để ảnh kiểm tra tham gia huấn luyện, và báo cáo cùng chỉ số. Nếu PatchCore cải thiện rõ rệt so với Isolation Forest trên đặc trưng thống kê, kết quả hỗ trợ vai trò của biểu diễn cục bộ sâu. Ngược lại, nếu đường cơ sở đạt kết quả tương đương, hệ thống đơn giản có thể được ưu tiên vì dễ giải thích và vận hành."
        ],
    },
    {
        "title": "2.5.15. Các họ phương pháp phát hiện bất thường",
        "paragraphs": [
            "Các phương pháp hiện đại có thể nhóm theo cách xây dựng tiêu chí bất thường. Nhóm dựa trên tái cấu trúc học cách phục hồi ảnh bình thường và dùng sai số phục hồi làm tín hiệu. Nhóm phân bố mô hình hóa đặc trưng bình thường bằng Gaussian, tâm hoặc mật độ. Nhóm ngân hàng bộ nhớ lưu patch tham chiếu và dùng láng giềng gần. Nhóm student–teacher so sánh đầu ra của mạng học sinh với mạng giáo viên cố định. Nhóm học khuyết tật tổng hợp tạo biến đổi giả để huấn luyện đầu phân biệt hoặc phân đoạn.",
            "Không có nhóm nào vượt trội trong mọi tình huống. Tái cấu trúc có thể vô tình phục hồi cả lỗi; ngân hàng bộ nhớ đòi hỏi dung lượng và tìm kiếm; mô hình phân bố phụ thuộc giả định; student–teacher cần quy trình huấn luyện ổn định; khuyết tật tổng hợp phụ thuộc mức tương đồng giữa lỗi giả và lỗi thật. Lựa chọn phải dựa trên quy mô dữ liệu, loại lỗi, yêu cầu định vị, tài nguyên và khả năng bảo trì.",
            "Trong chuyên đề, PatchCore được chọn làm phương pháp chính vì có thể sử dụng backbone tiền huấn luyện và không cần tối ưu gradient trên tập nhỏ. DRAEM và SuperSimpleNet đại diện cho hướng học từ lỗi tổng hợp, còn Isolation Forest là đường cơ sở thống kê. Việc đặt các phương pháp trong cùng hệ phân loại giúp lý giải vì sao kết quả của một mô hình không thể suy rộng trực tiếp sang mô hình khác: chúng sử dụng tín hiệu học và giả định về miền bình thường khác nhau."
        ],
        "table": [
            ["Họ phương pháp", "Tín hiệu bất thường", "Điểm cần lưu ý"],
            ["Tái cấu trúc", "Sai số ảnh/đặc trưng", "Có thể phục hồi cả lỗi"],
            ["Phân bố", "Độ lệch khỏi mô hình", "Phụ thuộc giả định phân bố"],
            ["Bộ nhớ", "Khoảng cách láng giềng", "Dung lượng và tốc độ tìm kiếm"],
            ["Student–teacher", "Sai khác đặc trưng", "Cần huấn luyện ổn định"],
            ["Lỗi tổng hợp", "Đầu phân biệt/phân đoạn", "Độ chân thực của lỗi giả"],
        ],
        "caption": "Bảng 2.5. So sánh các họ phương pháp phát hiện bất thường.",
    },
    {
        "title": "2.5.16. Mô hình tái cấu trúc: Autoencoder, VAE và GAN",
        "paragraphs": [
            "Autoencoder gồm bộ mã hóa đưa ảnh về biểu diễn tiềm ẩn và bộ giải mã khôi phục ảnh. Khi chỉ học trên mẫu bình thường với giới hạn năng lực phù hợp, mô hình được kỳ vọng tái cấu trúc tốt cấu trúc hợp lệ nhưng tạo sai số lớn tại vùng lỗi. Điểm có thể là sai số L1, L2, SSIM hoặc khoảng cách đặc trưng. Hạn chế cốt lõi là mạng có năng lực cao đôi khi học ánh xạ gần đồng nhất và phục hồi cả bất thường, làm giảm độ tương phản của bản đồ lỗi.",
            "Variational Autoencoder đặt phân bố xác suất cho biến tiềm ẩn và tối ưu đồng thời sai số tái cấu trúc với độ lệch KL [19]. Biểu diễn trơn hỗ trợ mô hình hóa miền bình thường nhưng ảnh tái cấu trúc có thể bị mờ, bất lợi cho lỗi chi tiết. GAN học bộ sinh thông qua cạnh tranh với bộ phân biệt [18]. AnoGAN tìm mã tiềm ẩn tạo ảnh gần quan sát rồi dùng phần dư và đặc trưng phân biệt để tính điểm [16], nhưng tối ưu mã cho từng ảnh làm suy luận chậm và vẫn phụ thuộc khả năng biểu diễn của bộ sinh.",
            "Các phương pháp tái cấu trúc cần được đánh giá không chỉ bằng chất lượng ảnh nhìn được. Một ảnh phục hồi đẹp có thể làm mất lỗi nhưng cũng có thể làm thay đổi vùng bình thường, tạo dương tính giả. So sánh trong không gian đặc trưng thường bền hơn pixel đối với dịch chuyển nhỏ. Đối với dữ liệu sản phẩm có nền, tái cấu trúc nền dễ chiếm phần lớn hàm mất mát; giới hạn vùng đối tượng hoặc dùng trọng số vùng quan tâm là cần thiết để tín hiệu lỗi nhỏ không bị pha loãng."
        ],
    },
    {
        "title": "2.5.17. Khuyết tật tổng hợp và phương pháp DRAEM",
        "paragraphs": [
            "Khi không có đủ ảnh lỗi thật, một hướng tiếp cận là tạo khuyết tật giả trên ảnh bình thường rồi huấn luyện mô hình nhận biết vùng bị thay đổi. Khuyết tật tổng hợp có thể được tạo bằng nhiễu Perlin, texture ngoài miền, phép cắt dán hoặc biến đổi màu. Mục tiêu không phải mô phỏng chính xác từng lỗi tương lai mà tạo nhiệm vụ học buộc mạng phân biệt cấu trúc bình thường với nhiễu có hình dạng và mức độ đa dạng.",
            "DRAEM kết hợp mạng tái cấu trúc và mạng phân biệt [5]. Ảnh bị làm lỗi giả được đưa qua mạng tái cấu trúc để phục hồi phiên bản gần bình thường; mạng phân biệt nhận ảnh gốc bị lỗi cùng ảnh phục hồi và dự đoán mặt nạ. Cách học phân biệt trực tiếp giúp tạo bản đồ sắc nét hơn nhiều autoencoder thuần túy. Tuy vậy, hiệu quả phụ thuộc phân bố lỗi tổng hợp. Nếu lỗi giả quá dễ, mạng học dấu hiệu nhân tạo; nếu quá khác lỗi thật, khả năng chuyển giao bị hạn chế.",
            "Thiết kế bộ tạo lỗi cần giữ tính hợp lệ của nền và chỉ thay đổi vùng có thể thuộc đối tượng. Tỷ lệ diện tích, độ trong suốt, texture và hình dạng nên được lấy từ khoảng rộng nhưng không phi thực tế. Quá trình huấn luyện cũng cần đủ ảnh bình thường để tránh ghi nhớ. Trong bối cảnh chỉ có 80 ảnh huấn luyện mỗi mã, phương pháp có tham số học phải được giám sát bằng đường cong mất mát và tập xác thực. Kết quả không hội tụ là thông tin quan trọng về điều kiện áp dụng, không nên bị loại khỏi báo cáo."
        ],
    },
    {
        "title": "2.5.18. Student–teacher và chưng cất tri thức",
        "paragraphs": [
            "Phương pháp student–teacher sử dụng một mạng giáo viên đã được tiền huấn luyện và giữ cố định. Mạng học sinh được tối ưu để bắt chước đặc trưng của giáo viên trên ảnh bình thường. Khi gặp vùng bất thường, giả thuyết là học sinh chưa được học để khớp đầu ra tại vùng đó, nên sai khác giữa hai mạng tăng. Công trình Uninformed Students khai thác nhiều học sinh và độ bất định để phát hiện, định vị bất thường [11].",
            "Tín hiệu sai khác có thể được tính ở nhiều tầng và nhiều tỷ lệ. Chuẩn hóa đặc trưng là cần thiết vì từng kênh có biên độ khác nhau. Nếu học sinh có năng lực quá lớn hoặc được tăng cường dữ liệu không phù hợp, nó có thể khái quát sang cả lỗi và làm giảm điểm. Nếu năng lực quá nhỏ, sai số trên vùng bình thường vẫn cao. Do đó, kiến trúc, hàm mất mát và lịch học tạo thành một cân bằng giữa bắt chước miền bình thường với duy trì độ nhạy ngoài miền.",
            "So với ngân hàng bộ nhớ, student–teacher nén kiến thức bình thường vào trọng số học sinh, có thể giảm chi phí lưu trữ và tìm kiếm. Đổi lại, nó cần huấn luyện và kiểm soát hội tụ. Khi dữ liệu mỗi mã ít, có thể huấn luyện chung nhiều mã nếu có cơ chế điều kiện hóa, nhưng điều đó làm miền bình thường rộng hơn. Chuyên đề ưu tiên PatchCore để tránh rủi ro tối ưu trên tập nhỏ; phần student–teacher vẫn cung cấp cơ sở đối chiếu cho hướng phát triển khi có thêm dữ liệu và tài nguyên huấn luyện."
        ],
    },
    {
        "title": "2.5.19. PaDiM và mô hình phân bố Gaussian theo vị trí",
        "paragraphs": [
            "PaDiM ghép đặc trưng từ nhiều tầng của CNN tiền huấn luyện và mô hình hóa phân bố patch bình thường tại từng vị trí không gian bằng Gaussian đa biến [10]. Với mỗi vị trí (i,j), phương pháp ước lượng vector trung bình μ_(i,j) và ma trận hiệp phương sai Σ_(i,j) từ ảnh huấn luyện. Patch kiểm tra được chấm bằng khoảng cách Mahalanobis đến phân bố tại đúng vị trí đó. Bản đồ khoảng cách sau nội suy trở thành bản đồ bất thường.",
            "Mô hình theo vị trí tận dụng bố cục ổn định: vùng góc trái chỉ được so sánh với cùng vùng trên ảnh bình thường, nhờ đó giảm nhầm lẫn giữa các bộ phận khác nhau. Tuy nhiên, nếu đối tượng dịch chuyển hoặc thay đổi tỷ lệ, patch hợp lệ rơi vào vị trí khác có thể bị chấm cao. Căn chỉnh ảnh hoặc thu nhận cố định vì thế đặc biệt quan trọng. PaDiM chọn ngẫu nhiên một số chiều đặc trưng để giảm chi phí ước lượng và nghịch đảo ma trận hiệp phương sai.",
            "So với PatchCore, PaDiM tóm tắt mỗi vị trí bằng tham số phân bố thay vì giữ tập mẫu. Điều này tiết kiệm tìm kiếm nhưng giả định phân bố gần Gaussian và liên kết mạnh với vị trí. PatchCore dùng ngân hàng chung nên linh hoạt hơn trước dịch chuyển cục bộ, song có thể so sánh nhầm các vùng giống nhau ở vị trí khác. Hai phương pháp minh họa đánh đổi giữa prior không gian, khả năng biểu diễn đa đỉnh và chi phí tính toán. Việc lựa chọn cần dựa trên mức ổn định bố cục của ảnh sản phẩm."
        ],
    },
    {
        "title": "2.5.20. Nguyên lý của PatchCore",
        "paragraphs": [
            "PatchCore hướng tới độ thu hồi cao trong phát hiện bất thường công nghiệp bằng cách lưu một tập đại diện các đặc trưng patch bình thường [4]. Ảnh huấn luyện đi qua backbone tiền huấn luyện; đặc trưng trung gian được căn chỉnh và tổng hợp thành vector patch. Toàn bộ vector tạo ngân hàng bộ nhớ ban đầu. Phương pháp không cập nhật trọng số backbone, nên quá trình “huấn luyện” chủ yếu là trích xuất đặc trưng và rút gọn ngân hàng.",
            "Khi suy luận, mỗi patch kiểm tra được tìm láng giềng gần nhất trong ngân hàng. Khoảng cách nhỏ cho thấy có cấu trúc bình thường tương tự; khoảng cách lớn tạo điểm bất thường cao. Điểm cấp ảnh thường dựa trên patch bất thường nhất và có thể được tái trọng số theo cấu trúc lân cận để giảm ảnh hưởng của một kết quả ghép không ổn định. Các điểm patch được sắp về lưới, nội suy và làm trơn để tạo bản đồ ở kích thước ảnh.",
            "Điểm mạnh của PatchCore là sử dụng trực tiếp độ đa dạng của ảnh bình thường, không ép phân bố vào một tâm hoặc Gaussian. Phương pháp phù hợp với tập nhỏ vì tận dụng đặc trưng tiền huấn luyện. Hạn chế gồm dung lượng ngân hàng, chi phí tìm kiếm, phụ thuộc vào backbone và khả năng báo động trước biến thiên bình thường chưa được quan sát. Ngoài ra, bản đồ định vị là hệ quả của lưới khoảng cách chứ không phải đầu ra phân đoạn được học bằng mặt nạ, nên đường biên chỉ mang tính xấp xỉ và cần hậu xử lý."
        ],
    },
    {
        "title": "2.5.21. Coreset và bài toán bao phủ ngân hàng đặc trưng",
        "paragraphs": [
            "Nếu lưu mọi patch của mọi ảnh, ngân hàng bộ nhớ tăng rất nhanh. Coreset là tập con được chọn để xấp xỉ độ bao phủ của tập đầy đủ. Ý tưởng k-center greedy chọn tuần tự điểm xa nhất so với tập đã chọn, nhằm giảm khoảng cách lớn nhất từ một điểm dữ liệu đến đại diện gần nhất. Cách tiếp cận core-set cũng được sử dụng trong học chủ động để bao phủ không gian đặc trưng [15] và được PatchCore áp dụng nhằm giữ hiệu năng với ngân hàng nhỏ hơn [4].",
            "Giả sử tập đặc trưng là F và coreset là C, mục tiêu trực quan là giảm max_(f∈F) min_(c∈C)||f−c||₂ dưới giới hạn kích thước |C|. Đây là bài toán khó tối ưu chính xác, nên thuật toán tham lam cung cấp nghiệm gần đúng. PatchCore có thể dùng phép chiếu ngẫu nhiên để giảm chiều trong giai đoạn chọn, sau đó lưu vector gốc của các điểm đại diện. Tỷ lệ coreset càng nhỏ thì suy luận càng nhanh nhưng nguy cơ bỏ mất cấu trúc bình thường hiếm càng cao.",
            "Tỷ lệ phù hợp không chỉ phụ thuộc số ảnh mà còn phụ thuộc độ đa dạng nội lớp và độ phân giải. Với bố cục gần cố định, nhiều patch lặp lại nên có thể rút gọn mạnh. Với vật liệu có hoa văn biến thiên, coreset quá nhỏ dễ làm tăng dương tính giả. Đánh giá nên báo cáo đồng thời AUROC, dung lượng ngân hàng và thời gian suy luận. Khi bổ sung ảnh bình thường mới, cần xây dựng lại hoặc cập nhật coreset có kiểm soát; chỉ nối thêm vector có thể làm ngân hàng mất cân bằng theo giai đoạn thu thập."
        ],
    },
    {
        "title": "2.5.22. Điểm cấp ảnh và bản đồ định vị bất thường",
        "paragraphs": [
            "Từ lưới khoảng cách patch A∈R^(h×w), hệ thống cần tạo hai đầu ra. Bản đồ định vị được nội suy đến H×W và có thể làm trơn Gaussian để giảm dạng ô lưới. Điểm cấp ảnh có thể là max(A), trung bình của nhóm điểm cao nhất hoặc một hàm tái trọng số. Max nhạy với lỗi nhỏ nhưng cũng nhạy với một patch nhiễu; trung bình toàn bản đồ ổn định hơn nhưng làm loãng lỗi nhỏ. Trung bình top-k là một lựa chọn trung gian cần được xác định trước khi đánh giá.",
            "Bản đồ điểm liên tục giữ nhiều thông tin hơn mặt nạ nhị phân. Khi thay đổi ngưỡng τ, mặt nạ B(x,y)=1[A(x,y)≥τ] thay đổi về diện tích và số vùng. Một ngưỡng dùng cho điểm cấp ảnh không nhất thiết phù hợp cho pixel vì hai phân bố khác nhau. Ngoài ra, điểm ảnh có tương quan không gian và bị ảnh hưởng bởi nội suy, nên không thể coi mỗi pixel là quan sát độc lập.",
            "Định vị bằng PatchCore nên được diễn giải như bản đồ vùng nghi ngờ. Trường tiếp nhận của patch khiến điểm cao có thể lan rộng hơn lỗi thật; làm trơn tiếp tục mở rộng biên. Nếu mục tiêu là khoanh vùng cho nhân viên kiểm tra, lớp phủ bán trong suốt và đường bao có thể hữu ích hơn mặt nạ tuyệt đối. Nếu mục tiêu là đo diện tích lỗi, cần mô hình phân đoạn được hiệu chỉnh và tập mặt nạ chuẩn đủ lớn. Chuyên đề tách rõ đánh giá cấp ảnh với đánh giá định vị để tránh suy luận quá mức từ một chỉ số duy nhất."
        ],
    },
    {
        "title": "2.5.23. Phân đoạn đối tượng và giới hạn vùng quan tâm",
        "paragraphs": [
            "Nền ảnh có thể tạo bất thường không liên quan đến chất lượng sản phẩm như bóng đổ, vật thể lạ hoặc thay đổi mép bàn. Mặt nạ đối tượng Ω cho phép chỉ giữ A_Ω(x,y)=A(x,y) nếu (x,y)∈Ω và loại điểm ngoài vùng. Cách này giảm dương tính giả nhưng chuyển độ tin cậy của hệ thống sang bước phân đoạn. Nếu mặt nạ bỏ sót một phần đối tượng, lỗi ở phần đó không còn cơ hội được phát hiện.",
            "Có thể tạo mặt nạ bằng kênh alpha, ngưỡng màu, mô hình phân đoạn hoặc kết hợp. Kênh alpha thật thường chính xác khi ảnh đã được tách nền; alpha giả được tạo từ ảnh RGB không mang bảo đảm tương tự. Các mạng phân đoạn ảnh nhị phân như IS-Net được thiết kế để tạo biên chính xác cho bài toán tách chủ thể [12], nhưng vẫn cần kiểm thử trên vật thể tối, trong suốt, phản xạ và có lỗ rỗng. Mọi mặt nạ nên được kiểm tra diện tích, số thành phần và tỷ lệ bao phủ như các tiêu chí chất lượng đầu vào.",
            "Giới hạn vùng quan tâm nên được áp dụng nhất quán cho cả hiệu chỉnh ngưỡng và suy luận. Nếu ngưỡng được tính từ toàn ảnh nhưng khi chạy chỉ xét tiền cảnh, phân bố điểm đã thay đổi. Trong chuyên đề, việc ưu tiên alpha thật và kiểm tra mặt nạ theo độ sáng xuất phát từ nguyên tắc này. Một cơ chế dự phòng hợp lý là phát hiện mặt nạ bất thường — quá nhỏ, quá lớn hoặc rỗng — rồi chuyển sang ảnh gốc hoặc yêu cầu kiểm tra thủ công, thay vì im lặng trả về kết quả có vẻ hợp lệ."
        ],
    },
    {
        "title": "2.5.24. Phân ngưỡng, hình thái học và thành phần liên thông",
        "paragraphs": [
            "Phân ngưỡng chuyển bản đồ liên tục thành quyết định. Ngưỡng cố định τ đơn giản và dễ triển khai nhưng giả định thang điểm ổn định giữa mã sản phẩm và thời gian. Ngưỡng theo phân vị chọn τ=Q_q(S_normal), với Q_q là phân vị q của điểm trên tập bình thường. Cách này kiểm soát gần đúng phần đuôi phân bố tham chiếu, song nếu tập hiệu chỉnh nhỏ thì ước lượng phân vị cao rất không ổn định. Ngưỡng cần được lưu cùng phiên bản mô hình và dữ liệu hiệu chỉnh.",
            "Sau ngưỡng, phép mở loại các điểm nhỏ, phép đóng lấp khe hở, còn giãn/co điều chỉnh biên. Các thao tác hình thái học dùng phần tử cấu trúc nên kích thước kernel phải tương ứng với độ phân giải ảnh và kích thước lỗi nhỏ nhất cần giữ. Kernel quá lớn có thể xóa vết xước thật hoặc nối các vùng không liên quan. Hậu xử lý không nên được tối ưu trực tiếp trên tập kiểm tra vì sẽ làm sai lệch đánh giá.",
            "Phân tích thành phần liên thông nhóm pixel kề nhau thành vùng. Giữ thành phần lớn nhất hữu ích khi kỳ vọng một lỗi chính, nhưng có thể loại nhiều lỗi nhỏ hợp lệ. Thay vào đó có thể lọc theo diện tích tối thiểu, tỷ lệ hình dạng hoặc khoảng cách tới biên đối tượng. Mỗi quy tắc đều chứa giả định nghiệp vụ và cần được ghi trong cấu hình. Để bảo toàn khả năng kiểm tra, hệ thống nên lưu cả bản đồ điểm gốc và mặt nạ sau hậu xử lý, giúp phân biệt lỗi của mô hình với lỗi do quy tắc nhị phân."
        ],
    },
    {
        "title": "2.5.25. Tổng hợp đa tỷ lệ và đánh đổi tính toán",
        "paragraphs": [
            "Khuyết tật có kích thước đa dạng nên một độ phân giải duy nhất khó tối ưu cho mọi trường hợp. Phân tích đa tỷ lệ chạy mô hình ở nhiều kích thước hoặc lấy đặc trưng từ nhiều tầng. Tỷ lệ thấp cung cấp ngữ cảnh và ổn định với nhiễu; tỷ lệ cao giữ chi tiết nhỏ. Sau khi đưa các bản đồ về cùng kích thước, có thể lấy trung bình, cực đại hoặc tổng có trọng số. Cực đại ưu tiên độ nhạy nhưng dễ giữ nhiễu; trung bình giảm nhiễu nhưng có thể làm yếu tín hiệu chỉ xuất hiện ở một tỷ lệ.",
            "Tăng độ phân giải làm số patch tăng gần theo diện tích. Nếu mỗi chiều tăng từ 224 lên 384, số vị trí có thể tăng khoảng (384/224)^2 trước khi xét stride, kéo theo thời gian trích xuất, tìm kiếm và bộ nhớ. Do đó, cải thiện định vị phải được cân đối với thời gian chu kỳ của dây chuyền. Có thể chỉ chạy tỷ lệ cao khi điểm ở tỷ lệ thấp nằm trong vùng chưa chắc chắn, tạo quy trình hai giai đoạn tiết kiệm tài nguyên.",
            "Tổng hợp đa tỷ lệ chỉ có ý nghĩa khi bản đồ được căn chỉnh đúng với ảnh gốc và cùng thang điểm. Nếu mỗi tỷ lệ có phân bố khoảng cách khác nhau, cần chuẩn hóa bằng thống kê của tập bình thường trước khi kết hợp. Trong nghiên cứu, nên so sánh từng tỷ lệ đơn lẻ với cấu hình tổng hợp và báo cáo chi phí bổ sung. Nếu chênh lệch chỉ số nhỏ, kết luận thận trọng là chưa có đủ bằng chứng cho lợi ích vận hành, thay vì mặc định cấu hình phức tạp hơn luôn tốt hơn."
        ],
    },
    {
        "title": "2.5.26. Chỉ số đánh giá và mối liên hệ với thiết kế nghiên cứu",
        "paragraphs": [
            "Ở cấp ảnh, ma trận nhầm lẫn gồm TP, FP, TN và FN tại một ngưỡng. Độ nhạy TPR=TP/(TP+FN) đo tỷ lệ ảnh lỗi được phát hiện; tỷ lệ dương giả FPR=FP/(FP+TN) đo phần ảnh bình thường bị cảnh báo. Đường ROC biểu diễn TPR theo FPR khi quét ngưỡng, còn AUROC là xác suất một mẫu lỗi được xếp điểm cao hơn một mẫu bình thường [13]. AUROC không xác định ngưỡng vận hành và có thể che khuất chi phí dương giả khi dữ liệu mất cân bằng, nên cần đi kèm kết quả tại ngưỡng được chọn.",
            "Ở cấp pixel, IoU=|P∩G|/|P∪G| và Dice=2|P∩G|/(|P|+|G|) so sánh mặt nạ dự đoán P với mặt nạ chuẩn G. Pixel AUROC coi từng pixel là quan sát nhưng dễ bị chi phối bởi số pixel nền lớn. PRO đo mức chồng lấp theo từng vùng lỗi trong một miền tỷ lệ dương giả, phù hợp hơn khi kích thước lỗi khác nhau. Pointing game chỉ kiểm tra điểm cao nhất có nằm trong vùng lỗi hay không, phản ánh khả năng chỉ điểm nhưng không đánh giá đường biên.",
            "Các lý thuyết trên dẫn trực tiếp đến thiết kế chuyên đề: backbone tiền huấn luyện và đặc trưng patch giải quyết dữ liệu ít; PatchCore và coreset mô hình hóa miền bình thường với chi phí có kiểm soát; mặt nạ đối tượng giảm nhiễu nền nhưng phải được kiểm định; đa tỷ lệ hỗ trợ lỗi nhỏ; ngưỡng và thành phần liên thông chuyển bản đồ điểm thành vùng cảnh báo. Việc báo cáo riêng cấp ảnh, định vị và chi phí tính toán giúp kết luận không vượt quá bằng chứng. Đồng thời, các phát hiện về vật thể tối và dương tính giả được xem là giới hạn có thể truy nguyên từ giả định tiền xử lý và hiệu chỉnh ngưỡng."
        ],
        "table": [
            ["Chỉ số", "Cấp đánh giá", "Ý nghĩa chính"],
            ["AUROC", "Ảnh hoặc pixel", "Khả năng xếp hạng qua mọi ngưỡng"],
            ["TPR/FPR", "Ảnh", "Hiệu năng tại ngưỡng vận hành"],
            ["IoU/Dice", "Pixel", "Mức chồng lấp mặt nạ"],
            ["PRO", "Vùng lỗi", "Bao phủ vùng theo dải dương giả"],
            ["Pointing game", "Vùng lỗi", "Điểm cao nhất có trúng vùng lỗi"],
        ],
        "caption": "Bảng 2.6. Các chỉ số đánh giá thường dùng.",
    },
]


REFERENCES = [
    "[7] A. Krizhevsky, I. Sutskever và G. E. Hinton, “ImageNet Classification with Deep Convolutional Neural Networks,” Advances in Neural Information Processing Systems, tập 25, 2012.",
    "[8] K. He, X. Zhang, S. Ren và J. Sun, “Deep Residual Learning for Image Recognition,” Proceedings of the IEEE Conference on Computer Vision and Pattern Recognition (CVPR), tr. 770–778, 2016.",
    "[9] L. Ruff, R. Vandermeulen, N. Görnitz, L. Deecke, S. A. Siddiqui, A. Binder, E. Müller và M. Kloft, “Deep One-Class Classification,” Proceedings of the 35th International Conference on Machine Learning, PMLR 80, tr. 4393–4402, 2018.",
    "[10] T. Defard, A. Setkov, A. Loesch và R. Audigier, “PaDiM: A Patch Distribution Modeling Framework for Anomaly Detection and Localization,” ICPR International Workshops and Challenges, Springer, tr. 475–489, 2021.",
    "[11] P. Bergmann, M. Fauser, D. Sattlegger và C. Steger, “Uninformed Students: Student–Teacher Anomaly Detection with Discriminative Latent Embeddings,” Proceedings of the IEEE/CVF Conference on Computer Vision and Pattern Recognition (CVPR), tr. 4183–4192, 2020.",
    "[12] X. Qin, H. Dai, X. Hu, D.-P. Fan, L. Shao và L. Van Gool, “Highly Accurate Dichotomous Image Segmentation,” European Conference on Computer Vision (ECCV), 2022.",
    "[13] T. Fawcett, “An Introduction to ROC Analysis,” Pattern Recognition Letters, tập 27, số 8, tr. 861–874, 2006.",
    "[14] P. Bergmann, K. Batzner, M. Fauser, D. Sattlegger và C. Steger, “Beyond Dents and Scratches: Logical Constraints in Unsupervised Anomaly Detection and Localization,” International Journal of Computer Vision, 2022.",
    "[15] O. Sener và S. Savarese, “Active Learning for Convolutional Neural Networks: A Core-Set Approach,” International Conference on Learning Representations (ICLR), 2018.",
    "[16] T. Schlegl, P. Seeböck, S. M. Waldstein, U. Schmidt-Erfurth và G. Langs, “Unsupervised Anomaly Detection with Generative Adversarial Networks to Guide Marker Discovery,” Information Processing in Medical Imaging, tr. 146–157, 2017.",
    "[17] J. Deng, W. Dong, R. Socher, L.-J. Li, K. Li và L. Fei-Fei, “ImageNet: A Large-Scale Hierarchical Image Database,” Proceedings of the IEEE Conference on Computer Vision and Pattern Recognition (CVPR), tr. 248–255, 2009.",
    "[18] I. Goodfellow và cộng sự, “Generative Adversarial Nets,” Advances in Neural Information Processing Systems, tập 27, 2014.",
    "[19] D. P. Kingma và M. Welling, “Auto-Encoding Variational Bayes,” International Conference on Learning Representations (ICLR), 2014.",
]


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_run(run, size=13, bold=False, italic=False):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def format_body(p):
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.first_line_indent = Cm(1.27)
    pf.line_spacing = 1.5
    pf.space_after = Pt(3)
    pf.space_before = Pt(0)
    pf.widow_control = True
    for r in p.runs:
        set_run(r)


def add_heading(doc, text, first=False):
    p = doc.add_paragraph(style="Heading 2")
    if first:
        p.add_run().add_break(WD_BREAK.PAGE)
    r = p.add_run(text)
    set_run(r, size=13, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0 if first else 0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    return p


def add_table(doc, rows, caption):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, row in enumerate(rows):
        for j, value in enumerate(row):
            cell = table.cell(i, j)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            set_run(r, size=11, bold=(i == 0))
            if i == 0:
                shade_cell(cell, "D9EAF7")
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(3)
    cp.paragraph_format.space_after = Pt(0)
    rr = cp.add_run(caption)
    set_run(rr, size=11, italic=True)


def build_block():
    block = Document()
    section = block.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)
    for i, item in enumerate(PAGES):
        add_heading(block, item["title"], first=True)
        for text in item["paragraphs"]:
            p = block.add_paragraph(text)
            format_body(p)
        if item.get("table"):
            add_table(block, item["table"], item["caption"])
    end = block.add_paragraph()
    end.add_run().add_break(WD_BREAK.PAGE)
    return block


def insert_before(target_paragraph, elements):
    for element in elements:
        target_paragraph._p.addprevious(deepcopy(element))


def ensure_styles(doc):
    styles = doc.styles
    if "Heading 2" not in [s.name for s in styles]:
        styles.add_style("Heading 2", WD_STYLE_TYPE.PARAGRAPH)
    h2 = styles["Heading 2"]
    h2.font.name = FONT
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    h2.font.size = Pt(13)
    h2.font.bold = True
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def append_references(doc, marker_text, marker_mode="exact"):
    marker = None
    for p in doc.paragraphs:
        t = p.text.strip()
        if (marker_mode == "exact" and t == marker_text) or (marker_mode == "starts" and t.startswith(marker_text)):
            marker = p
            break
    if marker is None:
        raise RuntimeError(f"Không tìm thấy vị trí chèn tài liệu tham khảo: {marker_text}")
    tmp = Document()
    for ref in REFERENCES:
        p = tmp.add_paragraph(ref)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(3)
        for r in p.runs:
            set_run(r, size=11)
    insert_before(marker, list(tmp.element.body)[:-1])


def update_document(path, chapter2_old, chapter2_new, chapter3_marker, refs_marker, refs_mode):
    doc = Document(path)
    if any(p.text.strip().startswith("2.5. Cơ sở lý thuyết") for p in doc.paragraphs):
        raise RuntimeError(f"Tệp đã có phần cơ sở lý thuyết: {path.name}")
    ensure_styles(doc)
    found_ch2 = False
    target = None
    for p in doc.paragraphs:
        t = p.text.strip()
        if t == chapter2_old:
            p.text = chapter2_new
            for r in p.runs:
                set_run(r, size=14 if path.name.startswith("IE400") else 13, bold=True)
            found_ch2 = True
        if t == chapter3_marker:
            target = p
    if not found_ch2 or target is None:
        raise RuntimeError(f"Không tìm đủ mốc trong {path.name}: chapter2={found_ch2}, chapter3={target is not None}")

    block = build_block()
    insert_before(target, list(block.element.body)[:-1])
    append_references(doc, refs_marker, refs_mode)
    doc.save(path)


def main():
    update_document(
        ROOT / "IE400.F3.CN2.CNTT_CDTN_NguyenThiHien_25410207.docx",
        "CHƯƠNG 2. TỔNG QUAN NGHIÊN CỨU",
        "CHƯƠNG 2. TỔNG QUAN NGHIÊN CỨU VÀ CƠ SỞ LÝ THUYẾT",
        "CHƯƠNG 3. DỮ LIỆU VÀ PHÁT BIỂU BÀI TOÁN",
        "PHỤ LỤC A. THỬ NGHIỆM SUPERSIMPLENET",
        "exact",
    )
    update_document(
        ROOT / "hienmain (1)_L.docx",
        "2. Các công trình nghiên cứu liên quan",
        "2. Các công trình nghiên cứu liên quan và cơ sở lý thuyết",
        "3. Thiết lập bộ dữ liệu và bài toán",
        "Hình 8:",
        "starts",
    )
    print(f"Đã bổ sung {len(PAGES)} trang nội dung lý thuyết và {len(REFERENCES)} tài liệu tham khảo vào mỗi tệp.")


if __name__ == "__main__":
    main()
