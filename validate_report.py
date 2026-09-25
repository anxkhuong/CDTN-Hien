from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader


docx_path = Path("output/Bao_cao_chuyen_de_tot_nghiep_50_trang.docx")
pdf_path = Path("tmp/Bao_cao_chuyen_de_tot_nghiep_final_check2.pdf")
doc = Document(docx_path)
text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
headings = [
    paragraph.text
    for paragraph in doc.paragraphs
    if paragraph.style and paragraph.style.name.startswith("Heading")
]
section = doc.sections[-1]

print(f"docx_bytes={docx_path.stat().st_size}")
print(f"pdf_pages={len(PdfReader(pdf_path).pages)}")
print(f"word_count={len(re.findall(r'\S+', text))}")
print(f"paragraphs={len(doc.paragraphs)}")
print(f"tables={len(doc.tables)}")
print(f"drawings={len(doc._element.xpath('.//w:drawing'))}")
print(f"headings={len(headings)}")
print(
    "body_margins_cm="
    f"{section.top_margin.cm:.1f},{section.bottom_margin.cm:.1f},"
    f"{section.left_margin.cm:.1f},{section.right_margin.cm:.1f}"
)
print(f"placeholder_fields={len(re.findall(r'\[[^\]]+\]', text))}")
print(f"toc_placeholder_present={'Cập nhật trường trong Word' in text}")
